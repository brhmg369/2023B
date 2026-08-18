"""Question 3 survey-line design for CUMCM 2023B.

The script designs north-south survey lines parallel to depth contours in the
simple rectangular slope area, verifies full coverage and adjacent overlap
constraints, and writes result3.xlsx.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
RESULT_XLSX = ROOT / "result3.xlsx"

NAUTICAL_MILE_M = 1852.0
THETA_DEG = 120.0
ALPHA_DEG = 1.5
CENTER_DEPTH_M = 110.0
EAST_WEST_WIDTH_NM = 4.0
NORTH_SOUTH_LENGTH_NM = 2.0
MIN_OVERLAP_RATE = 0.10
MAX_OVERLAP_RATE = 0.20
TARGET_OVERLAP_RATE = MIN_OVERLAP_RATE

X_MIN_M = -EAST_WEST_WIDTH_NM * NAUTICAL_MILE_M / 2.0
X_MAX_M = EAST_WEST_WIDTH_NM * NAUTICAL_MILE_M / 2.0
Y_MIN_M = -NORTH_SOUTH_LENGTH_NM * NAUTICAL_MILE_M / 2.0
Y_MAX_M = NORTH_SOUTH_LENGTH_NM * NAUTICAL_MILE_M / 2.0
LINE_LENGTH_M = Y_MAX_M - Y_MIN_M


@dataclass(frozen=True)
class SurveyLine:
    index: int
    x_m: float
    y_start_m: float
    y_end_m: float
    depth_m: float
    west_half_width_m: float
    east_half_width_m: float
    coverage_width_m: float
    west_edge_m: float
    east_edge_m: float
    spacing_from_previous_m: float | None
    overlap_width_m: float | None
    overlap_rate_pct: float | None


def depth_at(x_m: float, center_depth_m: float, alpha_rad: float) -> float:
    """Depth at east-west coordinate x; positive x is the shallow east side."""
    return center_depth_m - x_m * math.tan(alpha_rad)


def half_widths(depth_m: float, theta_rad: float, alpha_rad: float) -> tuple[float, float]:
    """Return horizontal west/east half-widths for a north-south survey line."""
    half_angle = theta_rad / 2.0
    common = depth_m * math.sin(half_angle) * math.cos(alpha_rad)
    west = common / math.cos(half_angle + alpha_rad)
    east = common / math.cos(half_angle - alpha_rad)
    return west, east


def west_edge(x_m: float, theta_rad: float, alpha_rad: float) -> float:
    west_half, _ = half_widths(depth_at(x_m, CENTER_DEPTH_M, alpha_rad), theta_rad, alpha_rad)
    return x_m - west_half


def east_edge(x_m: float, theta_rad: float, alpha_rad: float) -> float:
    _, east_half = half_widths(depth_at(x_m, CENTER_DEPTH_M, alpha_rad), theta_rad, alpha_rad)
    return x_m + east_half


def overlap_rate(previous_x_m: float, current_x_m: float, theta_rad: float, alpha_rad: float) -> float:
    """Adjacent overlap rate using the current line's horizontal swath width."""
    previous_depth = depth_at(previous_x_m, CENTER_DEPTH_M, alpha_rad)
    current_depth = depth_at(current_x_m, CENTER_DEPTH_M, alpha_rad)
    _, previous_east_half = half_widths(previous_depth, theta_rad, alpha_rad)
    current_west_half, current_east_half = half_widths(current_depth, theta_rad, alpha_rad)
    current_width = current_west_half + current_east_half
    overlap_width = previous_east_half + current_west_half - (current_x_m - previous_x_m)
    return overlap_width / current_width


def solve_first_line_x(theta_rad: float, alpha_rad: float) -> float:
    """Place the first line as far east as possible while covering the west edge."""
    lo = X_MIN_M - 2000.0
    hi = X_MAX_M
    for _ in range(100):
        mid = (lo + hi) / 2.0
        if west_edge(mid, theta_rad, alpha_rad) <= X_MIN_M:
            lo = mid
        else:
            hi = mid
    return lo


def solve_next_line_x(previous_x_m: float, theta_rad: float, alpha_rad: float) -> float:
    """Maximize the next x while keeping the overlap rate at least 10%."""
    lo = previous_x_m
    hi = max(X_MAX_M, previous_x_m + 1.0)
    while overlap_rate(previous_x_m, hi, theta_rad, alpha_rad) > TARGET_OVERLAP_RATE:
        hi += 1000.0
        if hi > X_MAX_M + 10000.0:
            raise RuntimeError("Failed to bracket the next survey line position.")

    for _ in range(100):
        mid = (lo + hi) / 2.0
        if overlap_rate(previous_x_m, mid, theta_rad, alpha_rad) >= TARGET_OVERLAP_RATE:
            lo = mid
        else:
            hi = mid
    return lo


def make_line(
    index: int,
    x_m: float,
    theta_rad: float,
    alpha_rad: float,
    previous: SurveyLine | None = None,
) -> SurveyLine:
    depth = depth_at(x_m, CENTER_DEPTH_M, alpha_rad)
    west_half, east_half = half_widths(depth, theta_rad, alpha_rad)
    spacing = None
    overlap_width = None
    overlap_rate_pct = None
    if previous is not None:
        spacing = x_m - previous.x_m
        overlap_width = previous.east_half_width_m + west_half - spacing
        overlap_rate_pct = overlap_width / (west_half + east_half) * 100.0

    return SurveyLine(
        index=index,
        x_m=x_m,
        y_start_m=Y_MIN_M,
        y_end_m=Y_MAX_M,
        depth_m=depth,
        west_half_width_m=west_half,
        east_half_width_m=east_half,
        coverage_width_m=west_half + east_half,
        west_edge_m=x_m - west_half,
        east_edge_m=x_m + east_half,
        spacing_from_previous_m=spacing,
        overlap_width_m=overlap_width,
        overlap_rate_pct=overlap_rate_pct,
    )


def design_lines() -> list[SurveyLine]:
    theta_rad = math.radians(THETA_DEG)
    alpha_rad = math.radians(ALPHA_DEG)

    lines: list[SurveyLine] = []
    x_m = solve_first_line_x(theta_rad, alpha_rad)
    lines.append(make_line(1, x_m, theta_rad, alpha_rad))

    while lines[-1].east_edge_m < X_MAX_M:
        x_m = solve_next_line_x(lines[-1].x_m, theta_rad, alpha_rad)
        lines.append(make_line(len(lines) + 1, x_m, theta_rad, alpha_rad, lines[-1]))

    return lines


def assert_flat_limit() -> None:
    theta_rad = math.radians(THETA_DEG)
    flat_west, flat_east = half_widths(CENTER_DEPTH_M, theta_rad, 0.0)
    expected = CENTER_DEPTH_M * math.tan(theta_rad / 2.0)
    if not math.isclose(flat_west, expected, rel_tol=1e-12, abs_tol=1e-12):
        raise AssertionError((flat_west, expected))
    if not math.isclose(flat_east, expected, rel_tol=1e-12, abs_tol=1e-12):
        raise AssertionError((flat_east, expected))


def max_east_edge_with_line_count(count: int) -> float:
    """Best east coverage achievable by count lines under the selected model."""
    if count < 1:
        raise ValueError("count must be positive")

    theta_rad = math.radians(THETA_DEG)
    alpha_rad = math.radians(ALPHA_DEG)
    x_m = solve_first_line_x(theta_rad, alpha_rad)
    for _ in range(count - 1):
        x_m = solve_next_line_x(x_m, theta_rad, alpha_rad)
    return east_edge(x_m, theta_rad, alpha_rad)


def validate_design(lines: list[SurveyLine]) -> dict[str, float]:
    if not lines:
        raise AssertionError("No survey lines generated.")

    tol = 1e-7
    if lines[0].west_edge_m > X_MIN_M + tol:
        raise AssertionError("The west boundary is not covered.")
    if lines[-1].east_edge_m < X_MAX_M - tol:
        raise AssertionError("The east boundary is not covered.")

    overlap_rates = [line.overlap_rate_pct for line in lines[1:]]
    for rate in overlap_rates:
        if rate is None:
            raise AssertionError("Missing overlap rate.")
        if rate < MIN_OVERLAP_RATE * 100.0 - 1e-6 or rate > MAX_OVERLAP_RATE * 100.0 + 1e-6:
            raise AssertionError(f"Overlap rate out of range: {rate}")

    previous_count_east_edge = max_east_edge_with_line_count(len(lines) - 1)
    if previous_count_east_edge >= X_MAX_M - tol:
        raise AssertionError("The line count is not minimal for the selected orientation.")

    return {
        "line_count": float(len(lines)),
        "total_length_m": len(lines) * LINE_LENGTH_M,
        "total_length_nm": len(lines) * NORTH_SOUTH_LENGTH_NM,
        "west_boundary_m": X_MIN_M,
        "east_boundary_m": X_MAX_M,
        "first_west_edge_m": lines[0].west_edge_m,
        "last_east_edge_m": lines[-1].east_edge_m,
        "east_excess_m": lines[-1].east_edge_m - X_MAX_M,
        "min_overlap_pct": min(overlap_rates),
        "max_overlap_pct": max(overlap_rates),
        "previous_count_east_edge_m": previous_count_east_edge,
        "previous_count_gap_m": X_MAX_M - previous_count_east_edge,
    }


def write_workbook(lines: list[SurveyLine], summary: dict[str, float]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "方案摘要"

    summary_rows = [
        ("测线方向", "南北向，平行等深线"),
        ("坐标约定", "x向东为正，y向北为正，原点为海域中心"),
        ("东西宽度/m", EAST_WEST_WIDTH_NM * NAUTICAL_MILE_M),
        ("南北长度/m", NORTH_SOUTH_LENGTH_NM * NAUTICAL_MILE_M),
        ("换能器开角/deg", THETA_DEG),
        ("坡度/deg", ALPHA_DEG),
        ("中心水深/m", CENTER_DEPTH_M),
        ("重叠率下限/%", MIN_OVERLAP_RATE * 100.0),
        ("重叠率上限/%", MAX_OVERLAP_RATE * 100.0),
        ("设计测线数/条", int(summary["line_count"])),
        ("测线总长度/m", summary["total_length_m"]),
        ("测线总长度/NM", summary["total_length_nm"]),
        ("最小相邻重叠率/%", summary["min_overlap_pct"]),
        ("最大相邻重叠率/%", summary["max_overlap_pct"]),
        ("西边界/m", summary["west_boundary_m"]),
        ("首条测线西覆盖边界/m", summary["first_west_edge_m"]),
        ("东边界/m", summary["east_boundary_m"]),
        ("末条测线东覆盖边界/m", summary["last_east_edge_m"]),
        ("东侧超出覆盖/m", summary["east_excess_m"]),
        ("若少一条线的最东覆盖边界/m", summary["previous_count_east_edge_m"]),
        ("若少一条线的东侧缺口/m", summary["previous_count_gap_m"]),
    ]
    ws.append(["项目", "数值"])
    for row in summary_rows:
        ws.append(list(row))

    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 34

    detail = wb.create_sheet("测线坐标与覆盖")
    detail_headers = [
        "line_id",
        "x/m",
        "x/NM",
        "y_start/m",
        "y_start/NM",
        "y_end/m",
        "y_end/NM",
        "line_length/m",
        "D/m",
        "west_half_width/m",
        "east_half_width/m",
        "W_h/m",
        "west_edge/m",
        "east_edge/m",
        "spacing_from_previous/m",
        "overlap_width/m",
        "overlap_rate/%",
    ]
    detail.append(detail_headers)
    for line in lines:
        detail.append(
            [
                line.index,
                line.x_m,
                line.x_m / NAUTICAL_MILE_M,
                line.y_start_m,
                line.y_start_m / NAUTICAL_MILE_M,
                line.y_end_m,
                line.y_end_m / NAUTICAL_MILE_M,
                LINE_LENGTH_M,
                line.depth_m,
                line.west_half_width_m,
                line.east_half_width_m,
                line.coverage_width_m,
                line.west_edge_m,
                line.east_edge_m,
                line.spacing_from_previous_m,
                line.overlap_width_m,
                line.overlap_rate_pct,
            ]
        )

    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column in detail.columns:
        detail.column_dimensions[column[0].column_letter].width = 20
        for cell in column:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    checks = wb.create_sheet("相邻重叠校验")
    checks.append(
        [
            "pair",
            "previous_line",
            "current_line",
            "spacing/m",
            "overlap_width/m",
            "current_W_h/m",
            "overlap_rate/%",
            "within_10_to_20_pct",
        ]
    )
    for line in lines[1:]:
        checks.append(
            [
                f"{line.index - 1}-{line.index}",
                line.index - 1,
                line.index,
                line.spacing_from_previous_m,
                line.overlap_width_m,
                line.coverage_width_m,
                line.overlap_rate_pct,
                MIN_OVERLAP_RATE * 100.0 <= line.overlap_rate_pct <= MAX_OVERLAP_RATE * 100.0,
            ]
        )

    for cell in checks[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FCE4D6")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in checks.columns:
        checks.column_dimensions[column[0].column_letter].width = 22
        for cell in column:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    wb.save(RESULT_XLSX)


def print_summary(lines: list[SurveyLine], summary: dict[str, float]) -> None:
    print("Question 3 survey-line design")
    print(f"line_count: {len(lines)}")
    print(f"total_length_m: {summary['total_length_m']:.2f}")
    print(f"total_length_nm: {summary['total_length_nm']:.2f}")
    print(f"overlap_pct_range: {summary['min_overlap_pct']:.2f} - {summary['max_overlap_pct']:.2f}")
    print(f"first_west_edge_m: {summary['first_west_edge_m']:.2f}")
    print(f"last_east_edge_m: {summary['last_east_edge_m']:.2f}")
    print(f"previous_count_gap_m: {summary['previous_count_gap_m']:.2f}")
    print("line x positions in meters:")
    print([round(line.x_m, 2) for line in lines])


def main() -> None:
    assert_flat_limit()
    lines = design_lines()
    summary = validate_design(lines)
    write_workbook(lines, summary)
    print_summary(lines, summary)
    print(f"Wrote {RESULT_XLSX}")


if __name__ == "__main__":
    main()
