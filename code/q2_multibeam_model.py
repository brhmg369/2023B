"""Question 2 three-dimensional multibeam swath-width model.

The script computes Table 2 for CUMCM 2023B and writes result2.xlsx.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
RESULT_XLSX = ROOT / "result2.xlsx"

NAUTICAL_MILE_M = 1852.0
THETA_DEG = 120.0
ALPHA_DEG = 1.5
CENTER_DEPTH_M = 120.0
DISTANCES_NM = [0.0, 0.3, 0.6, 0.9, 1.2, 1.5, 1.8, 2.1]
BETAS_DEG = [0, 45, 90, 135, 180, 225, 270, 315]


@dataclass(frozen=True)
class CellResult:
    beta_deg: int
    distance_nm: float
    distance_m: float
    depth_m: float
    effective_slope_deg: float
    coverage_width_m: float


def depth_at(distance_nm: float, beta_rad: float, center_depth_m: float, alpha_rad: float) -> float:
    """Depth along the survey-line direction from the sea-area center.

    The positive horizontal projection of the seabed normal is taken as the
    deeper direction. With this convention beta=0 moves toward deeper water,
    beta=180 moves toward shallower water, and beta=90/270 stays on a contour.
    """
    distance_m = distance_nm * NAUTICAL_MILE_M
    return center_depth_m + distance_m * math.tan(alpha_rad) * math.cos(beta_rad)


def effective_cross_slope(alpha_rad: float, beta_rad: float) -> float:
    """Slope angle in the vertical plane perpendicular to the survey line."""
    return math.atan(math.tan(alpha_rad) * math.sin(beta_rad))


def coverage_width(depth_m: float, theta_rad: float, effective_slope_rad: float) -> float:
    """Horizontal swath width from the two-dimensional slope formula."""
    half_angle = theta_rad / 2.0
    return (
        depth_m
        * math.sin(half_angle)
        * math.cos(effective_slope_rad)
        * (
            1.0 / math.cos(half_angle + effective_slope_rad)
            + 1.0 / math.cos(half_angle - effective_slope_rad)
        )
    )


def compute_results() -> list[CellResult]:
    theta_rad = math.radians(THETA_DEG)
    alpha_rad = math.radians(ALPHA_DEG)
    results: list[CellResult] = []

    for beta_deg in BETAS_DEG:
        beta_rad = math.radians(beta_deg)
        effective_slope = effective_cross_slope(alpha_rad, beta_rad)
        for distance_nm in DISTANCES_NM:
            depth = depth_at(distance_nm, beta_rad, CENTER_DEPTH_M, alpha_rad)
            width = coverage_width(depth, theta_rad, effective_slope)
            results.append(
                CellResult(
                    beta_deg=beta_deg,
                    distance_nm=distance_nm,
                    distance_m=distance_nm * NAUTICAL_MILE_M,
                    depth_m=depth,
                    effective_slope_deg=math.degrees(effective_slope),
                    coverage_width_m=width,
                )
            )

    return results


def result_map(results: list[CellResult]) -> dict[tuple[int, float], CellResult]:
    return {(item.beta_deg, item.distance_nm): item for item in results}


def assert_model_limits(results: list[CellResult]) -> None:
    lookup = result_map(results)
    theta_rad = math.radians(THETA_DEG)

    # At beta=0 the cross-track section is parallel to a contour, so the model
    # must reduce to the flat-bottom width at the local depth.
    for distance_nm in DISTANCES_NM:
        item = lookup[(0, distance_nm)]
        expected = 2.0 * item.depth_m * math.tan(theta_rad / 2.0)
        if not math.isclose(item.coverage_width_m, expected, rel_tol=1e-12, abs_tol=1e-12):
            raise AssertionError((item, expected))

    # At beta=90 or beta=270 the ship moves along a contour, so local depth
    # stays equal to the center depth for all table distances.
    for beta_deg in (90, 270):
        for distance_nm in DISTANCES_NM:
            item = lookup[(beta_deg, distance_nm)]
            if not math.isclose(item.depth_m, CENTER_DEPTH_M, rel_tol=1e-12, abs_tol=1e-12):
                raise AssertionError(item)


def write_workbook(results: list[CellResult]) -> None:
    lookup = result_map(results)

    wb = Workbook()
    ws = wb.active
    ws.title = "问题2结果"

    ws["A1"] = "覆盖宽度/m"
    ws["B1"] = "测量船距海域中心点处的距离/海里"
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=len(DISTANCES_NM) + 1)
    ws["A2"] = "测线方向夹角/deg"

    for col_idx, distance_nm in enumerate(DISTANCES_NM, start=2):
        ws.cell(2, col_idx, distance_nm)

    for row_idx, beta_deg in enumerate(BETAS_DEG, start=3):
        ws.cell(row_idx, 1, beta_deg)
        for col_idx, distance_nm in enumerate(DISTANCES_NM, start=2):
            ws.cell(row_idx, col_idx, round(lookup[(beta_deg, distance_nm)].coverage_width_m, 2))

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=3, max_row=len(BETAS_DEG) + 2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    ws.column_dimensions["A"].width = 20
    for col_idx in range(2, len(DISTANCES_NM) + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 13

    detail = wb.create_sheet("计算明细")
    detail_headers = [
        "beta/deg",
        "distance/NM",
        "distance/m",
        "D/m",
        "alpha_eff/deg",
        "coverage_width/m",
    ]
    detail.append(detail_headers)
    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in results:
        detail.append(
            [
                item.beta_deg,
                item.distance_nm,
                item.distance_m,
                item.depth_m,
                item.effective_slope_deg,
                item.coverage_width_m,
            ]
        )

    for column in detail.columns:
        detail.column_dimensions[column[0].column_letter].width = 18
        for cell in column:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    wb.save(RESULT_XLSX)


def print_table(results: list[CellResult]) -> None:
    lookup = result_map(results)
    print("Table 2 results")
    print(["覆盖宽度/m"] + DISTANCES_NM)
    for beta_deg in BETAS_DEG:
        print([beta_deg] + [round(lookup[(beta_deg, distance)].coverage_width_m, 2) for distance in DISTANCES_NM])


def main() -> None:
    results = compute_results()
    assert_model_limits(results)
    write_workbook(results)
    print_table(results)
    print(f"Wrote {RESULT_XLSX}")


if __name__ == "__main__":
    main()
