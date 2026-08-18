"""Question 4 real-bathymetry survey-line design for CUMCM 2023B.

The script reads the single-beam depth grid from attachment.xlsx, estimates
local cross-track slopes, designs north-south segmented survey lines by
horizontal bands, evaluates coverage and excessive overlap on the supplied
grid, and writes result4.xlsx.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = ROOT / "附件.xlsx"
RESULT_XLSX = ROOT / "result4.xlsx"

NAUTICAL_MILE_M = 1852.0
THETA_DEG = 120.0
SELECTED_BAND_HEIGHT_NM = 0.5
CANDIDATE_BAND_HEIGHTS_NM = [5.0, 1.0, 0.5, 0.25]
OVERLAP_LIMIT = 0.20


@dataclass(frozen=True)
class BathymetryGrid:
    x_nm: np.ndarray
    y_nm: np.ndarray
    depth_m: np.ndarray
    grad_x: np.ndarray
    grad_y: np.ndarray

    @property
    def x_m(self) -> np.ndarray:
        return self.x_nm * NAUTICAL_MILE_M

    @property
    def y_m(self) -> np.ndarray:
        return self.y_nm * NAUTICAL_MILE_M

    @property
    def dx_m(self) -> float:
        return float((self.x_nm[1] - self.x_nm[0]) * NAUTICAL_MILE_M)

    @property
    def dy_m(self) -> float:
        return float((self.y_nm[1] - self.y_nm[0]) * NAUTICAL_MILE_M)

    @property
    def x_min_m(self) -> float:
        return float(self.x_m[0])

    @property
    def x_max_m(self) -> float:
        return float(self.x_m[-1])

    @property
    def y_min_m(self) -> float:
        return float(self.y_m[0])

    @property
    def y_max_m(self) -> float:
        return float(self.y_m[-1])

    @property
    def area_m2(self) -> float:
        return (self.x_max_m - self.x_min_m) * (self.y_max_m - self.y_min_m)


@dataclass(frozen=True)
class Band:
    index: int
    y_start_m: float
    y_end_m: float
    row_indices: np.ndarray

    @property
    def height_m(self) -> float:
        return self.y_end_m - self.y_start_m


@dataclass(frozen=True)
class Segment:
    segment_id: int
    band_index: int
    line_index_in_band: int
    x_m: float
    y_start_m: float
    y_end_m: float
    length_m: float
    mean_depth_m: float
    min_depth_m: float
    max_depth_m: float
    mean_width_m: float
    min_width_m: float
    max_width_m: float
    west_edge_min_m: float
    east_edge_max_m: float


@dataclass(frozen=True)
class Scheme:
    name: str
    band_height_nm: float
    bands: list[Band]
    segments_by_band: list[list[Segment]]

    @property
    def segments(self) -> list[Segment]:
        return [segment for band_segments in self.segments_by_band for segment in band_segments]


@dataclass(frozen=True)
class SchemeMetrics:
    name: str
    band_height_nm: float
    band_count: int
    segment_count: int
    total_length_m: float
    total_length_nm: float
    missed_area_pct: float
    covered_area_pct: float
    excessive_overlap_length_m: float
    excessive_overlap_length_nm: float
    max_overlap_rate_pct: float
    mean_overlap_rate_pct: float


@dataclass(frozen=True)
class OverlapRecord:
    scheme_name: str
    band_index: int
    previous_segment_id: int
    current_segment_id: int
    spacing_m: float
    mean_overlap_rate_pct: float
    max_overlap_rate_pct: float
    excessive_overlap_length_m: float


def read_bathymetry() -> BathymetryGrid:
    wb = load_workbook(INPUT_XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    x_nm = np.array([float(ws.cell(2, col).value) for col in range(3, ws.max_column + 1)], dtype=float)
    y_nm = np.array([float(ws.cell(row, 2).value) for row in range(3, ws.max_row + 1)], dtype=float)
    depth = np.array(
        [
            [float(ws.cell(row, col).value) for col in range(3, ws.max_column + 1)]
            for row in range(3, ws.max_row + 1)
        ],
        dtype=float,
    )

    dy_m = (y_nm[1] - y_nm[0]) * NAUTICAL_MILE_M
    dx_m = (x_nm[1] - x_nm[0]) * NAUTICAL_MILE_M
    grad_y, grad_x = np.gradient(depth, dy_m, dx_m)

    if depth.shape != (len(y_nm), len(x_nm)):
        raise AssertionError("Depth matrix shape does not match coordinate axes.")
    if np.isnan(depth).any():
        raise AssertionError("Depth matrix contains missing values.")

    return BathymetryGrid(x_nm=x_nm, y_nm=y_nm, depth_m=depth, grad_x=grad_x, grad_y=grad_y)


def interpolate_x(values: np.ndarray, grid: BathymetryGrid, x_m: float, rows: np.ndarray) -> np.ndarray:
    x_axis = grid.x_m
    if x_m <= x_axis[0]:
        return values[rows, 0]
    if x_m >= x_axis[-1]:
        return values[rows, -1]

    left_index = int(np.searchsorted(x_axis, x_m) - 1)
    left_index = max(0, min(left_index, len(x_axis) - 2))
    weight = (x_m - x_axis[left_index]) / (x_axis[left_index + 1] - x_axis[left_index])
    return (1.0 - weight) * values[rows, left_index] + weight * values[rows, left_index + 1]


def half_widths(grid: BathymetryGrid, x_m: float, rows: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Horizontal west/east half-widths for north-south line segments.

    Depth is positive downward and x is positive eastward. In the local linear
    cross-track section D(x+u)=D(x)+g_x*u, the west/east intersections give
    B_w=D/(cot(phi)+g_x), B_e=D/(cot(phi)-g_x).
    """
    half_angle = math.radians(THETA_DEG) / 2.0
    cot_half_angle = 1.0 / math.tan(half_angle)
    depth = interpolate_x(grid.depth_m, grid, x_m, rows)
    grad_x = interpolate_x(grid.grad_x, grid, x_m, rows)
    west_denominator = cot_half_angle + grad_x
    east_denominator = cot_half_angle - grad_x

    if np.min(west_denominator) <= 0 or np.min(east_denominator) <= 0:
        raise AssertionError("Local slope is outside the valid swath-width range.")

    return depth / west_denominator, depth / east_denominator


def make_bands(grid: BathymetryGrid, band_height_nm: float) -> list[Band]:
    band_count = int(round((grid.y_nm[-1] - grid.y_nm[0]) / band_height_nm))
    if not math.isclose(band_count * band_height_nm, grid.y_nm[-1] - grid.y_nm[0], abs_tol=1e-9):
        raise ValueError("Band height must divide the north-south range exactly.")

    bands: list[Band] = []
    for index in range(1, band_count + 1):
        y_start_nm = grid.y_nm[0] + (index - 1) * band_height_nm
        y_end_nm = grid.y_nm[0] + index * band_height_nm
        y_start_m = y_start_nm * NAUTICAL_MILE_M
        y_end_m = y_end_nm * NAUTICAL_MILE_M
        if index < band_count:
            rows = np.where((grid.y_m >= y_start_m - 1e-9) & (grid.y_m < y_end_m - 1e-9))[0]
        else:
            rows = np.where((grid.y_m >= y_start_m - 1e-9) & (grid.y_m <= y_end_m + 1e-9))[0]
        if len(rows) == 0:
            raise AssertionError("Empty band generated.")
        bands.append(Band(index=index, y_start_m=y_start_m, y_end_m=y_end_m, row_indices=rows))
    return bands


def solve_first_x(grid: BathymetryGrid, rows: np.ndarray) -> float:
    lo = grid.x_min_m - 2000.0
    hi = grid.x_max_m
    for _ in range(80):
        mid = (lo + hi) / 2.0
        west_half, _ = half_widths(grid, mid, rows)
        if float(np.max(mid - west_half)) <= grid.x_min_m:
            lo = mid
        else:
            hi = mid
    return lo


def solve_next_x(grid: BathymetryGrid, previous_x_m: float, rows: np.ndarray) -> float:
    _, previous_east_half = half_widths(grid, previous_x_m, rows)
    lo = previous_x_m
    hi = grid.x_max_m
    for _ in range(80):
        mid = (lo + hi) / 2.0
        current_west_half, _ = half_widths(grid, mid, rows)
        no_gap = float(np.min(previous_east_half + current_west_half - (mid - previous_x_m))) >= 0.0
        if no_gap:
            lo = mid
        else:
            hi = mid
    return lo


def summarize_segment(
    grid: BathymetryGrid,
    segment_id: int,
    band: Band,
    line_index_in_band: int,
    x_m: float,
) -> Segment:
    rows = band.row_indices
    depth = interpolate_x(grid.depth_m, grid, x_m, rows)
    west_half, east_half = half_widths(grid, x_m, rows)
    width = west_half + east_half
    return Segment(
        segment_id=segment_id,
        band_index=band.index,
        line_index_in_band=line_index_in_band,
        x_m=x_m,
        y_start_m=band.y_start_m,
        y_end_m=band.y_end_m,
        length_m=band.height_m,
        mean_depth_m=float(np.mean(depth)),
        min_depth_m=float(np.min(depth)),
        max_depth_m=float(np.max(depth)),
        mean_width_m=float(np.mean(width)),
        min_width_m=float(np.min(width)),
        max_width_m=float(np.max(width)),
        west_edge_min_m=float(np.min(x_m - west_half)),
        east_edge_max_m=float(np.max(x_m + east_half)),
    )


def design_band(grid: BathymetryGrid, band: Band, next_segment_id: int) -> tuple[list[Segment], int]:
    rows = band.row_indices
    x_positions = [solve_first_x(grid, rows)]
    while True:
        _, east_half = half_widths(grid, x_positions[-1], rows)
        if float(np.min(x_positions[-1] + east_half)) >= grid.x_max_m - 1e-7:
            break
        next_x = solve_next_x(grid, x_positions[-1], rows)
        if next_x <= x_positions[-1] + 1e-7:
            raise RuntimeError(f"Line placement stalled in band {band.index}.")
        x_positions.append(next_x)
        if len(x_positions) > 1000:
            raise RuntimeError(f"Too many segments in band {band.index}.")

    segments: list[Segment] = []
    segment_id = next_segment_id
    for local_index, x_m in enumerate(x_positions, start=1):
        segments.append(summarize_segment(grid, segment_id, band, local_index, x_m))
        segment_id += 1
    return segments, segment_id


def design_scheme(grid: BathymetryGrid, band_height_nm: float) -> Scheme:
    name = "全域统一布线" if band_height_nm == 5.0 else f"{band_height_nm:g} NM分带布线"
    if math.isclose(band_height_nm, SELECTED_BAND_HEIGHT_NM):
        name += "（推荐）"

    bands = make_bands(grid, band_height_nm)
    segments_by_band: list[list[Segment]] = []
    next_segment_id = 1
    for band in bands:
        segments, next_segment_id = design_band(grid, band, next_segment_id)
        segments_by_band.append(segments)
    return Scheme(name=name, band_height_nm=band_height_nm, bands=bands, segments_by_band=segments_by_band)


def evaluate_scheme(grid: BathymetryGrid, scheme: Scheme) -> tuple[SchemeMetrics, list[OverlapRecord]]:
    covered = np.zeros(grid.depth_m.shape, dtype=bool)
    overlap_rates: list[float] = []
    excessive_overlap_length_m = 0.0
    overlap_records: list[OverlapRecord] = []

    for band, segments in zip(scheme.bands, scheme.segments_by_band):
        rows = band.row_indices
        row_x = grid.x_m[None, :]

        for segment in segments:
            west_half, east_half = half_widths(grid, segment.x_m, rows)
            covered[np.ix_(rows, np.arange(len(grid.x_m)))] |= (
                (row_x >= (segment.x_m - west_half)[:, None])
                & (row_x <= (segment.x_m + east_half)[:, None])
            )

        for previous, current in zip(segments[:-1], segments[1:]):
            _, previous_east_half = half_widths(grid, previous.x_m, rows)
            current_west_half, current_east_half = half_widths(grid, current.x_m, rows)
            current_width = current_west_half + current_east_half
            spacing = current.x_m - previous.x_m
            rate = (previous_east_half + current_west_half - spacing) / current_width
            overlap_rates.extend(rate.tolist())
            excessive_length = band.height_m * float(np.mean(rate > OVERLAP_LIMIT))
            excessive_overlap_length_m += excessive_length
            overlap_records.append(
                OverlapRecord(
                    scheme_name=scheme.name,
                    band_index=band.index,
                    previous_segment_id=previous.segment_id,
                    current_segment_id=current.segment_id,
                    spacing_m=spacing,
                    mean_overlap_rate_pct=float(np.mean(rate) * 100.0),
                    max_overlap_rate_pct=float(np.max(rate) * 100.0),
                    excessive_overlap_length_m=excessive_length,
                )
            )

    missed_area_pct = float((1.0 - np.mean(covered)) * 100.0)
    total_length_m = float(sum(segment.length_m for segment in scheme.segments))
    max_overlap_rate_pct = float(max(overlap_rates) * 100.0 if overlap_rates else 0.0)
    mean_overlap_rate_pct = float(np.mean(overlap_rates) * 100.0 if overlap_rates else 0.0)
    return (
        SchemeMetrics(
            name=scheme.name,
            band_height_nm=scheme.band_height_nm,
            band_count=len(scheme.bands),
            segment_count=len(scheme.segments),
            total_length_m=total_length_m,
            total_length_nm=total_length_m / NAUTICAL_MILE_M,
            missed_area_pct=missed_area_pct,
            covered_area_pct=100.0 - missed_area_pct,
            excessive_overlap_length_m=float(excessive_overlap_length_m),
            excessive_overlap_length_nm=float(excessive_overlap_length_m / NAUTICAL_MILE_M),
            max_overlap_rate_pct=max_overlap_rate_pct,
            mean_overlap_rate_pct=mean_overlap_rate_pct,
        ),
        overlap_records,
    )


def assert_data_contract(grid: BathymetryGrid) -> None:
    if grid.depth_m.shape != (251, 201):
        raise AssertionError(f"Unexpected depth grid shape: {grid.depth_m.shape}")
    if not math.isclose(grid.x_nm[0], 0.0) or not math.isclose(grid.x_nm[-1], 4.0):
        raise AssertionError("Unexpected x coordinate range.")
    if not math.isclose(grid.y_nm[0], 0.0) or not math.isclose(grid.y_nm[-1], 5.0):
        raise AssertionError("Unexpected y coordinate range.")
    if not math.isclose(grid.x_nm[1] - grid.x_nm[0], 0.02, abs_tol=1e-12):
        raise AssertionError("Unexpected x coordinate step.")
    if not math.isclose(grid.y_nm[1] - grid.y_nm[0], 0.02, abs_tol=1e-12):
        raise AssertionError("Unexpected y coordinate step.")


def write_workbook(
    grid: BathymetryGrid,
    schemes: list[Scheme],
    metrics: list[SchemeMetrics],
    selected_scheme: Scheme,
    selected_metrics: SchemeMetrics,
    selected_overlap_records: list[OverlapRecord],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "方案摘要"

    summary_rows = [
        ("推荐方案", selected_metrics.name),
        ("测线方向", "南北向线段，按0.5 NM水平带分区递推"),
        ("坐标单位", "附件坐标为海里，程序内部换算为米"),
        ("海域尺寸", "东西4 NM，南北5 NM"),
        ("换能器开角/deg", THETA_DEG),
        ("水深最小值/m", float(np.min(grid.depth_m))),
        ("水深最大值/m", float(np.max(grid.depth_m))),
        ("水深平均值/m", float(np.mean(grid.depth_m))),
        ("分带数", selected_metrics.band_count),
        ("测线线段数", selected_metrics.segment_count),
        ("测线总长度/m", selected_metrics.total_length_m),
        ("测线总长度/NM", selected_metrics.total_length_nm),
        ("漏测海区占比/%", selected_metrics.missed_area_pct),
        ("覆盖海区占比/%", selected_metrics.covered_area_pct),
        ("重叠率超过20%部分总长度/m", selected_metrics.excessive_overlap_length_m),
        ("重叠率超过20%部分总长度/NM", selected_metrics.excessive_overlap_length_nm),
        ("最大相邻重叠率/%", selected_metrics.max_overlap_rate_pct),
        ("平均相邻重叠率/%", selected_metrics.mean_overlap_rate_pct),
    ]
    ws.append(["项目", "数值"])
    for row in summary_rows:
        ws.append(list(row))
    style_sheet(ws, header_fill="D9EAF7", widths=[34, 38])

    comparison = wb.create_sheet("候选方案比较")
    comparison.append(
        [
            "方案",
            "分带高度/NM",
            "分带数",
            "测线线段数",
            "总长度/m",
            "总长度/NM",
            "漏测率/%",
            "覆盖率/%",
            "超20%重叠长度/m",
            "最大重叠率/%",
            "平均重叠率/%",
        ]
    )
    for item in metrics:
        comparison.append(
            [
                item.name,
                item.band_height_nm,
                item.band_count,
                item.segment_count,
                item.total_length_m,
                item.total_length_nm,
                item.missed_area_pct,
                item.covered_area_pct,
                item.excessive_overlap_length_m,
                item.max_overlap_rate_pct,
                item.mean_overlap_rate_pct,
            ]
        )
    style_sheet(comparison, header_fill="E2F0D9")

    bands_sheet = wb.create_sheet("推荐方案分带")
    bands_sheet.append(
        [
            "band_id",
            "y_start/NM",
            "y_end/NM",
            "line_segment_count",
            "band_length/m",
            "band_total_length/m",
            "first_x/NM",
            "last_x/NM",
        ]
    )
    for band, segments in zip(selected_scheme.bands, selected_scheme.segments_by_band):
        bands_sheet.append(
            [
                band.index,
                band.y_start_m / NAUTICAL_MILE_M,
                band.y_end_m / NAUTICAL_MILE_M,
                len(segments),
                band.height_m,
                len(segments) * band.height_m,
                segments[0].x_m / NAUTICAL_MILE_M,
                segments[-1].x_m / NAUTICAL_MILE_M,
            ]
        )
    style_sheet(bands_sheet, header_fill="FCE4D6")

    segments_sheet = wb.create_sheet("推荐测线坐标")
    segments_sheet.append(
        [
            "segment_id",
            "band_id",
            "line_id_in_band",
            "x/NM",
            "x/m",
            "y_start/NM",
            "y_end/NM",
            "segment_length/m",
            "mean_depth/m",
            "min_depth/m",
            "max_depth/m",
            "mean_W_h/m",
            "min_W_h/m",
            "max_W_h/m",
            "west_edge_min/m",
            "east_edge_max/m",
        ]
    )
    for segment in selected_scheme.segments:
        segments_sheet.append(
            [
                segment.segment_id,
                segment.band_index,
                segment.line_index_in_band,
                segment.x_m / NAUTICAL_MILE_M,
                segment.x_m,
                segment.y_start_m / NAUTICAL_MILE_M,
                segment.y_end_m / NAUTICAL_MILE_M,
                segment.length_m,
                segment.mean_depth_m,
                segment.min_depth_m,
                segment.max_depth_m,
                segment.mean_width_m,
                segment.min_width_m,
                segment.max_width_m,
                segment.west_edge_min_m,
                segment.east_edge_max_m,
            ]
        )
    style_sheet(segments_sheet, header_fill="DDEBF7")

    overlap_sheet = wb.create_sheet("推荐重叠校验")
    overlap_sheet.append(
        [
            "band_id",
            "previous_segment_id",
            "current_segment_id",
            "spacing/m",
            "mean_overlap_rate/%",
            "max_overlap_rate/%",
            "over20_length/m",
        ]
    )
    for record in selected_overlap_records:
        overlap_sheet.append(
            [
                record.band_index,
                record.previous_segment_id,
                record.current_segment_id,
                record.spacing_m,
                record.mean_overlap_rate_pct,
                record.max_overlap_rate_pct,
                record.excessive_overlap_length_m,
            ]
        )
    style_sheet(overlap_sheet, header_fill="FFF2CC")

    audit = wb.create_sheet("数据审计")
    audit.append(["项目", "数值"])
    audit_rows = [
        ("数据文件", str(INPUT_XLSX.name)),
        ("横向坐标数量", len(grid.x_nm)),
        ("纵向坐标数量", len(grid.y_nm)),
        ("水深点数量", int(grid.depth_m.size)),
        ("横向范围/NM", f"{grid.x_nm[0]:.2f}~{grid.x_nm[-1]:.2f}"),
        ("纵向范围/NM", f"{grid.y_nm[0]:.2f}~{grid.y_nm[-1]:.2f}"),
        ("坐标步长/NM", grid.x_nm[1] - grid.x_nm[0]),
        ("水深最小/m", float(np.min(grid.depth_m))),
        ("水深最大/m", float(np.max(grid.depth_m))),
        ("水深平均/m", float(np.mean(grid.depth_m))),
        ("gx最小", float(np.min(grid.grad_x))),
        ("gx最大", float(np.max(grid.grad_x))),
        ("gy最小", float(np.min(grid.grad_y))),
        ("gy最大", float(np.max(grid.grad_y))),
    ]
    for row in audit_rows:
        audit.append(list(row))
    style_sheet(audit, header_fill="EADCF8", widths=[30, 40])

    wb.save(RESULT_XLSX)


def style_sheet(ws, header_fill: str, widths: list[int] | None = None) -> None:
    fill = PatternFill("solid", fgColor=header_fill)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    if widths is not None:
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + idx)].width = width
    else:
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 20


def print_summary(metrics: list[SchemeMetrics], selected: SchemeMetrics) -> None:
    print("Question 4 candidate schemes")
    for index, item in enumerate(metrics, start=1):
        print(
            f"scheme_{index}: band_height_nm={item.band_height_nm:g}, "
            f"segments={item.segment_count}, "
            f"length={item.total_length_m:.2f} m, "
            f"miss={item.missed_area_pct:.4f}%, "
            f"over20={item.excessive_overlap_length_m:.2f} m"
        )
    print("Selected scheme")
    print(f"band_height_nm: {selected.band_height_nm:g}")
    print(f"total_length_m: {selected.total_length_m:.2f}")
    print(f"missed_area_pct: {selected.missed_area_pct:.4f}")
    print(f"excessive_overlap_length_m: {selected.excessive_overlap_length_m:.2f}")
    print(f"Wrote {RESULT_XLSX}")


def main() -> None:
    grid = read_bathymetry()
    assert_data_contract(grid)

    schemes = [design_scheme(grid, band_height) for band_height in CANDIDATE_BAND_HEIGHTS_NM]
    evaluated = [evaluate_scheme(grid, scheme) for scheme in schemes]
    metrics = [item[0] for item in evaluated]
    selected_index = CANDIDATE_BAND_HEIGHTS_NM.index(SELECTED_BAND_HEIGHT_NM)
    selected_scheme = schemes[selected_index]
    selected_metrics, selected_overlap_records = evaluated[selected_index]

    write_workbook(grid, schemes, metrics, selected_scheme, selected_metrics, selected_overlap_records)
    print_summary(metrics, selected_metrics)


if __name__ == "__main__":
    main()
