"""Generate the high-end figure pack for CUMCM 2023B.

The figures are built from the same formulas and result workbooks used by
the four modeling scripts. Each figure is saved as PDF, SVG, and PNG so the
paper writer can use vector output in LaTeX and raster previews for review.
"""

from __future__ import annotations

import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.cm import ScalarMappable
from matplotlib.colors import LinearSegmentedColormap, Normalize
from matplotlib.patches import Arc, Polygon, Rectangle


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
STYLE = FIG / "style" / "cumcm.mplstyle"
plt.style.use(STYLE)

NM = 1852.0
THETA_DEG = 120.0
ALPHA_DEG = 1.5

INK = "#2C3237"
MUTED = "#6F7880"
GRID = "#D8DEE3"
NAVY = "#2C4866"
BLUE = "#3B7893"
TEAL = "#5A9A8A"
AMBER = "#B99B4C"
CLAY = "#8B7A66"
WATER = "#E3F1F5"
SLOPE = "#E9E5DA"
ACCENT = "#A84D3B"

DEPTH_CMAP = LinearSegmentedColormap.from_list(
    "paper_depth",
    ["#F1E6C8", "#C7D7C7", "#7EA9AE", "#456F87", "#2F3D55"],
)
WIDTH_CMAP = LinearSegmentedColormap.from_list(
    "paper_width",
    ["#343F61", "#4C7D9A", "#7EA99B", "#C4AE71"],
)

plt.rcParams.update(
    {
        "axes.facecolor": "white",
        "figure.facecolor": "white",
        "savefig.facecolor": "white",
        "font.family": "sans-serif",
        "font.sans-serif": ["Microsoft YaHei", "Arial", "DejaVu Sans", "SimHei"],
        "mathtext.fontset": "dejavusans",
        "axes.edgecolor": INK,
        "axes.labelcolor": INK,
        "xtick.color": INK,
        "ytick.color": INK,
        "text.color": INK,
        "grid.color": GRID,
        "grid.alpha": 0.55,
        "grid.linewidth": 0.55,
        "legend.frameon": False,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
        "font.size": 9,
        "axes.titlesize": 10,
        "axes.labelsize": 9,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "legend.fontsize": 8,
        "axes.titleweight": "normal",
    }
)


def depth_at(x_m: float | np.ndarray, center_depth_m: float, alpha_rad: float) -> float | np.ndarray:
    return center_depth_m - x_m * np.tan(alpha_rad)


def half_widths(depth_m: float, theta_rad: float, alpha_rad: float) -> tuple[float, float]:
    half_angle = theta_rad / 2.0
    common = depth_m * math.sin(half_angle) * math.cos(alpha_rad)
    left = common / math.cos(half_angle + alpha_rad)
    right = common / math.cos(half_angle - alpha_rad)
    return left, right


def q2_coverage(r_nm: float | np.ndarray, beta_deg: float | np.ndarray, center_depth_m: float = 120.0):
    alpha_rad = np.deg2rad(ALPHA_DEG)
    theta_rad = np.deg2rad(THETA_DEG)
    beta_rad = np.deg2rad(beta_deg)
    s = np.asarray(r_nm) * NM
    depth = center_depth_m + s * np.tan(alpha_rad) * np.cos(beta_rad)
    alpha_eff = np.arctan(np.tan(alpha_rad) * np.sin(beta_rad))
    half_angle = theta_rad / 2.0
    return (
        depth
        * np.sin(half_angle)
        * np.cos(alpha_eff)
        * (1.0 / np.cos(half_angle + alpha_eff) + 1.0 / np.cos(half_angle - alpha_eff))
    )


def read_result3() -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    wb = openpyxl.load_workbook(ROOT / "result3.xlsx", data_only=True)
    ws = wb["测线坐标与覆盖"]
    x, spacing, eta, west, east = [], [], [], [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is None:
            continue
        x.append(float(row[1]))
        west.append(float(row[12]))
        east.append(float(row[13]))
        spacing.append(float(row[14]) if row[14] is not None else np.nan)
        eta.append(float(row[16]) if row[16] is not None else np.nan)
    wb.close()
    return np.array(x), np.array(spacing), np.array(eta), np.array(west), np.array(east)


def read_result4_comparison() -> dict[str, np.ndarray | list[str]]:
    wb = openpyxl.load_workbook(ROOT / "result4.xlsx", data_only=True)
    ws = wb["候选方案比较"]
    short_names = ["全域 5.00 NM", "1.00 NM", "0.50 NM", "0.25 NM"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows.append(row)
    data = {
        "names": short_names[: len(rows)],
        "band_height": np.array([float(r[1]) for r in rows]),
        "segments": np.array([float(r[3]) for r in rows]),
        "length": np.array([float(r[4]) for r in rows]),
        "over20": np.array([float(r[8]) for r in rows]),
        "max_eta": np.array([float(r[9]) for r in rows]),
        "avg_eta": np.array([float(r[10]) for r in rows]),
    }
    wb.close()
    return data


def read_result4_segments() -> np.ndarray:
    wb = openpyxl.load_workbook(ROOT / "result4.xlsx", data_only=True)
    ws = wb["推荐测线坐标"]
    segs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[3] is None:
            continue
        segs.append((float(row[3]), float(row[5]), float(row[6])))
    wb.close()
    return np.array(segs)


def read_bathymetry() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    wb = openpyxl.load_workbook(ROOT / "附件.xlsx", data_only=True)
    ws = wb[wb.sheetnames[0]]
    x = np.array([float(ws.cell(2, c).value) for c in range(3, ws.max_column + 1)])
    y = np.array([float(ws.cell(r, 2).value) for r in range(3, ws.max_row + 1)])
    z = np.array(
        [[float(ws.cell(r, c).value) for c in range(3, ws.max_column + 1)] for r in range(3, ws.max_row + 1)]
    )
    wb.close()
    return x, y, z


def save(fig, q: str, name: str) -> None:
    out = FIG / q
    out.mkdir(parents=True, exist_ok=True)
    for suffix in ("pdf", "svg"):
        fig.savefig(out / f"{name}.{suffix}", bbox_inches="tight", pad_inches=0.04)
    png = out / f"{name}.png"
    try:
        fig.savefig(png, dpi=300, bbox_inches="tight", pad_inches=0.04)
    except OSError:
        tmp_png = out / f"{name}.tmp.png"
        fig.savefig(tmp_png, dpi=300, bbox_inches="tight", pad_inches=0.04)
        tmp_png.replace(png)
    plt.close(fig)
    print(f"Wrote {out / name}.pdf / .svg / .png")


def clean_axis(ax, grid: bool = True) -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    if grid:
        ax.grid(True, zorder=0)
    else:
        ax.grid(False)


def arrow2d(ax, start, end, color=INK, lw=1.2, text: str | None = None, text_offset=(0, 0), **kwargs):
    ax.annotate(
        "",
        xy=end,
        xytext=start,
        arrowprops=dict(arrowstyle="->", color=color, lw=lw, shrinkA=0, shrinkB=0, **kwargs),
    )
    if text:
        ax.text((start[0] + end[0]) / 2 + text_offset[0], (start[1] + end[1]) / 2 + text_offset[1], text)


def double_arrow(ax, x0, x1, y, text, color=INK, text_y=0.0):
    ax.annotate(
        "",
        xy=(x1, y),
        xytext=(x0, y),
        arrowprops=dict(arrowstyle="<->", color=color, lw=1.05, shrinkA=0, shrinkB=0),
    )
    ax.text((x0 + x1) / 2, y + text_y, text, ha="center", va="bottom", color=color, fontsize=10)


def fig_q1_geometry() -> None:
    alpha = math.radians(ALPHA_DEG)
    theta = math.radians(THETA_DEG)
    d0 = 70.0
    x0 = 70.0
    d = float(depth_at(x0, d0, alpha))
    bl, br = half_widths(d, theta, alpha)
    xl, xr = x0 - bl, x0 + br
    x_min, x_max = xl - 85, xr + 85
    xs = np.linspace(x_min, x_max, 300)
    seabed = -depth_at(xs, d0, alpha)

    fig, ax = plt.subplots(figsize=(7.6, 4.15))
    ax.set_axis_off()
    ax.set_xlim(x_min, x_max)
    ax.set_ylim(min(seabed) - 18, 28)

    water_poly = np.column_stack([np.r_[x_min, x_max, xs[::-1]], np.r_[0, 0, seabed[::-1]]])
    bed_poly = np.column_stack(
        [np.r_[x_min, x_max, x_max, x_min], np.r_[seabed[0], seabed[-1], min(seabed) - 28, min(seabed) - 28]]
    )
    ax.add_patch(Polygon(water_poly, closed=True, facecolor=WATER, edgecolor="none", alpha=0.82, zorder=0))
    ax.add_patch(Polygon(bed_poly, closed=True, facecolor=SLOPE, edgecolor="none", alpha=0.95, zorder=0))
    ax.plot(xs, seabed, color=CLAY, lw=2.0, zorder=2)
    ax.plot([x_min, x_max], [0, 0], color=BLUE, lw=1.8, zorder=2)

    y_l = -float(depth_at(xl, d0, alpha))
    y_r = -float(depth_at(xr, d0, alpha))
    ax.plot([x0, xl], [0, y_l], color=ACCENT, lw=1.6, zorder=3)
    ax.plot([x0, xr], [0, y_r], color=ACCENT, lw=1.6, zorder=3)
    ax.fill([x0, xl, xr], [0, y_l, y_r], color=AMBER, alpha=0.15, zorder=1)
    ax.plot([x0, x0], [0, -d], color=MUTED, ls="--", lw=1.0, zorder=3)
    ax.plot([xl, xl], [0, y_l], color=MUTED, ls=":", lw=0.95, zorder=3)
    ax.plot([xr, xr], [0, y_r], color=MUTED, ls=":", lw=0.95, zorder=3)
    ax.scatter([x0, xl, xr], [0, y_l, y_r], s=[38, 24, 24], color=[BLUE, ACCENT, ACCENT], zorder=4)

    double_arrow(ax, xl, xr, 12, r"$W=B_L+B_R$（水平投影）", BLUE, text_y=2.5)
    double_arrow(ax, xl, x0, 5, r"$B_L$", TEAL, text_y=1.5)
    double_arrow(ax, x0, xr, 5, r"$B_R$", TEAL, text_y=1.5)
    ax.annotate(
        "",
        xy=(x0 + 18, -d),
        xytext=(x0 + 18, 0),
        arrowprops=dict(arrowstyle="<->", color=MUTED, lw=1.0),
    )
    ax.text(x0 + 25, -d / 2, r"$D(x)$", color=MUTED, va="center")

    ax.add_patch(Arc((x0, 0), 54, 54, theta1=238, theta2=302, color=ACCENT, lw=1.0))
    ax.text(x0 + 17, -22, r"$\theta$", color=ACCENT, ha="center")
    arc_x = x_min + 70
    arc_y = float(np.interp(arc_x, xs, seabed))
    ax.add_patch(Arc((arc_x, arc_y), 72, 34, angle=-ALPHA_DEG, theta1=0, theta2=15, color=CLAY, lw=1.0))
    ax.text(arc_x + 48, arc_y + 4, r"$\alpha$", color=CLAY)

    ax.text(x_min + 8, 4.5, "海平面", color=BLUE, fontsize=10)
    ax.text(x_max - 98, float(np.interp(x_max - 98, xs, seabed)) - 9, "斜坡海底", color=CLAY, fontsize=10)
    ax.text(x0 + 6, 4.2, "测线位置", color=BLUE, fontsize=10)
    save(fig, "q1", "fig_q1_geometry")


def fig_q1_result() -> None:
    theta = math.radians(THETA_DEG)
    alpha = math.radians(ALPHA_DEG)
    d0 = 70.0
    xs = np.linspace(-800, 800, 321)
    depth = np.array([depth_at(x, d0, alpha) for x in xs])
    width = np.empty_like(depth)
    for i, d in enumerate(depth):
        bl, br = half_widths(float(d), theta, alpha)
        width[i] = bl + br

    table_x = np.arange(-800, 801, 200)
    eta = []
    for i in range(1, len(table_x)):
        prev_d = float(depth_at(table_x[i - 1], d0, alpha))
        cur_d = float(depth_at(table_x[i], d0, alpha))
        _, prev_br = half_widths(prev_d, theta, alpha)
        cur_bl, cur_br = half_widths(cur_d, theta, alpha)
        spacing = table_x[i] - table_x[i - 1]
        eta.append((prev_br + cur_bl - spacing) / (cur_bl + cur_br) * 100.0)
    eta = np.array(eta)

    fig, axes = plt.subplots(1, 3, figsize=(11.6, 3.35), constrained_layout=True)
    axes[0].plot(xs, depth, color=BLUE, lw=2.2)
    axes[0].fill_between(xs, depth, depth.min() - 4, color=BLUE, alpha=0.12)
    axes[0].set_xlabel("测线位置 x / m")
    axes[0].set_ylabel("水深 D / m")
    axes[0].set_title("水深随位置变化")

    axes[1].plot(xs, width, color=AMBER, lw=2.2)
    axes[1].fill_between(xs, width, width.min() - 10, color=AMBER, alpha=0.16)
    axes[1].set_xlabel("测线位置 x / m")
    axes[1].set_ylabel("覆盖宽度 W / m")
    axes[1].set_title("水平覆盖宽度")

    axes[2].axhspan(10, 20, color=TEAL, alpha=0.12, lw=0)
    axes[2].plot(table_x[1:], eta, marker="o", color=TEAL, lw=2.0)
    axes[2].axhline(0, color=MUTED, ls=":", lw=1.0)
    axes[2].axhline(10, color=TEAL, ls="--", lw=1.0)
    axes[2].axhline(20, color=TEAL, ls="--", lw=1.0)
    axes[2].set_xlabel("当前测线位置 x / m")
    axes[2].set_ylabel("重叠率 eta / %")
    axes[2].set_title("固定间距下的重叠变化")
    axes[2].text(255, 16.5, "10%-20% 期望区间", color=TEAL, fontsize=9)
    axes[2].text(610, -7.2, "浅水侧出现漏测", color=ACCENT, fontsize=9)

    for ax in axes:
        clean_axis(ax)
    save(fig, "q1", "fig_q1_result")


def fig_q2_3d_geometry() -> None:
    beta = math.radians(135.0)
    heading = np.array([math.cos(beta), math.sin(beta), 0.0])
    cross = np.array([-math.sin(beta), math.cos(beta), 0.0])

    def bed_z(x_coord: float | np.ndarray) -> float | np.ndarray:
        return -0.44 - 0.18 * x_coord

    def project(points: np.ndarray) -> np.ndarray:
        pts = np.asarray(points, dtype=float)
        return np.stack((pts[..., 0] - 0.55 * pts[..., 1], 0.34 * pts[..., 1] + pts[..., 2]), axis=-1)

    def poly3d(ax, points, facecolor, edgecolor, lw=0.9, alpha=1.0, zorder=1):
        ax.add_patch(
            Polygon(
                project(np.array(points)),
                closed=True,
                facecolor=facecolor,
                edgecolor=edgecolor,
                lw=lw,
                alpha=alpha,
                joinstyle="miter",
                zorder=zorder,
            )
        )

    def line3d(ax, points, color=INK, lw=1.0, ls="-", alpha=1.0, zorder=4):
        pts = project(np.array(points))
        ax.plot(pts[:, 0], pts[:, 1], color=color, lw=lw, ls=ls, alpha=alpha, zorder=zorder)

    def arrow3d(ax, start, end, color=INK, lw=1.0, zorder=6):
        p0, p1 = project(np.array([start, end]))
        ax.annotate(
            "",
            xy=p1,
            xytext=p0,
            arrowprops=dict(arrowstyle="->", color=color, lw=lw, shrinkA=0, shrinkB=0),
            zorder=zorder,
        )

    def label3d(ax, point, text, color=INK, dx=0.0, dy=0.0, size=9, ha="left", va="center"):
        p = project(np.array(point))
        ax.text(p[0] + dx, p[1] + dy, text, color=color, fontsize=size, ha=ha, va=va, zorder=8)

    sea = np.array([[-1.12, -0.82, 0.0], [1.12, -0.82, 0.0], [1.12, 0.82, 0.0], [-1.12, 0.82, 0.0]])
    bed = np.array([[p[0], p[1], bed_z(p[0])] for p in sea])
    u0, u1 = -0.58, 0.58
    scan_top0 = u0 * cross + np.array([0, 0, 0.035])
    scan_top1 = u1 * cross + np.array([0, 0, 0.035])
    scan_bot1 = u1 * cross + np.array([0, 0, bed_z(u1 * cross[0])])
    scan_bot0 = u0 * cross + np.array([0, 0, bed_z(u0 * cross[0])])
    scan = np.vstack([scan_top0, scan_top1, scan_bot1, scan_bot0])

    fig, ax = plt.subplots(figsize=(7.0, 4.8))
    ax.set_axis_off()
    ax.set_aspect("equal")

    poly3d(ax, bed, SLOPE, "#A99F91", lw=0.9, alpha=0.98, zorder=1)
    for yy in (-0.44, 0.0, 0.44):
        line3d(ax, [[-1.12, yy, bed_z(-1.12)], [1.12, yy, bed_z(1.12)]], "#B7AEA2", lw=0.45, alpha=0.8, zorder=2)
    for xx in (-0.56, 0.0, 0.56):
        line3d(ax, [[xx, -0.82, bed_z(xx)], [xx, 0.82, bed_z(xx)]], "#B7AEA2", lw=0.42, alpha=0.62, zorder=2)
    poly3d(ax, sea, WATER, BLUE, lw=0.95, alpha=0.56, zorder=3)
    poly3d(ax, scan, "#EFE1B7", AMBER, lw=1.0, alpha=0.66, zorder=4)

    line3d(ax, [scan_top0, scan_top1], AMBER, lw=1.05, zorder=5)
    line3d(ax, [scan_bot0, scan_bot1], AMBER, lw=1.05, zorder=5)
    line3d(ax, [scan_top0, scan_bot0], AMBER, lw=0.72, ls="--", alpha=0.78, zorder=5)
    line3d(ax, [scan_top1, scan_bot1], AMBER, lw=0.72, ls="--", alpha=0.78, zorder=5)

    route0 = -0.74 * heading + np.array([0, 0, 0.055])
    route1 = 0.74 * heading + np.array([0, 0, 0.055])
    line3d(ax, [route0, route1], NAVY, lw=1.7, zorder=6)
    arrow3d(ax, -0.18 * heading + np.array([0, 0, 0.065]), 0.42 * heading + np.array([0, 0, 0.065]), NAVY, lw=1.1)

    origin = np.array([-0.78, -0.56, 0.075])
    arrow3d(ax, origin, origin + np.array([0.46, 0.0, 0.0]), ACCENT, lw=1.15)
    arrow3d(ax, origin, origin + 0.56 * heading, NAVY, lw=1.15)
    arc_t = np.linspace(0, beta, 64)
    arc = origin + np.column_stack([0.24 * np.cos(arc_t), 0.24 * np.sin(arc_t), np.zeros_like(arc_t)])
    line3d(ax, arc, INK, lw=0.9, zorder=7)

    beam_l = scan_bot0
    beam_r = scan_bot1
    sonar = np.array([0.0, 0.0, 0.06])
    line3d(ax, [sonar, beam_l], ACCENT, lw=0.92, alpha=0.82, zorder=6)
    line3d(ax, [sonar, beam_r], ACCENT, lw=0.92, alpha=0.82, zorder=6)
    line3d(ax, [sonar, np.array([0.0, 0.0, bed_z(0.0)])], MUTED, lw=0.82, ls="--", alpha=0.9, zorder=6)

    label3d(ax, [-1.02, 0.78, 0.0], "海平面", BLUE, dx=-0.04, dy=0.03, size=9)
    label3d(ax, [0.66, 0.74, bed_z(0.66)], "斜坡海底", CLAY, dx=0.02, dy=-0.04, size=9)
    label3d(ax, route1, "航向", NAVY, dx=0.02, dy=0.04, size=9)
    label3d(ax, scan_bot1, "扫描平面", AMBER, dx=0.03, dy=-0.01, size=9)
    label3d(ax, origin + np.array([0.25, 0.02, 0.0]), r"$\beta$", INK, dx=0.0, dy=0.0, size=10)
    label3d(ax, origin + np.array([0.46, 0.0, 0.0]), "坡面法向水平投影", ACCENT, dx=0.03, dy=-0.01, size=8)

    all_pts = np.vstack([sea, bed, scan, route0, route1, origin, origin + np.array([0.46, 0, 0]), origin + 0.56 * heading])
    p = project(all_pts)
    ax.set_xlim(p[:, 0].min() - 0.22, p[:, 0].max() + 0.24)
    ax.set_ylim(p[:, 1].min() - 0.16, p[:, 1].max() + 0.18)
    save(fig, "q2", "fig_q2_3d_geometry")


def fig_q2_heatmap() -> None:
    betas = np.linspace(0, 315, 127)
    rs = np.linspace(0, 2.1, 141)
    r_grid, beta_grid = np.meshgrid(rs, betas)
    width = q2_coverage(r_grid, beta_grid)

    fig, ax = plt.subplots(figsize=(7.3, 4.35), constrained_layout=True)
    levels = np.linspace(np.nanmin(width), np.nanmax(width), 16)
    im = ax.contourf(r_grid, beta_grid, width, levels=levels, cmap=WIDTH_CMAP)
    cs = ax.contour(r_grid, beta_grid, width, levels=np.arange(100, 801, 100), colors="white", linewidths=0.48, alpha=0.72)
    ax.clabel(cs, fmt="%d", fontsize=7, colors="white")
    cb = fig.colorbar(im, ax=ax, shrink=0.93, pad=0.015)
    cb.set_label("覆盖宽度 W / m")
    cb.set_ticks(np.arange(100, 701, 100))
    ax.set_xlabel("测量船距中心点距离 r / 海里")
    ax.set_ylabel(r"测线方向夹角 $\beta$ / deg")
    ax.set_title(r"$W(r,\beta)$ 的二维等值投影")
    for beta in (0, 90, 180, 270):
        ax.axhline(beta, color="white", lw=0.75, alpha=0.55)
    clean_axis(ax, grid=False)
    save(fig, "q2", "fig_q2_heatmap")


def fig_q2_width_surface() -> None:
    rs = np.linspace(0, 2.1, 58)
    betas = np.linspace(0, 315, 76)
    r_grid, beta_grid = np.meshgrid(rs, betas)
    width = q2_coverage(r_grid, beta_grid)

    fig = plt.figure(figsize=(7.6, 5.35))
    ax = fig.add_subplot(111, projection="3d")
    surf = ax.plot_surface(r_grid, beta_grid, width, cmap=WIDTH_CMAP, linewidth=0, antialiased=True, alpha=0.94)
    ax.contour(r_grid, beta_grid, width, zdir="z", offset=40, levels=10, cmap=WIDTH_CMAP, linewidths=0.66, alpha=0.72)
    cb = fig.colorbar(surf, ax=ax, shrink=0.62, pad=0.08, label="覆盖宽度 W / m")
    cb.set_ticks(np.arange(100, 701, 100))
    ax.set_xlabel("r / 海里")
    ax.set_ylabel(r"$\beta$ / deg")
    ax.set_zlabel("W / m")
    ax.set_zlim(40, 790)
    ax.set_title(r"$W(r,\beta)$ 覆盖宽度曲面")
    ax.view_init(elev=28, azim=-132)
    ax.set_box_aspect((1.25, 1.25, 0.72))
    ax.xaxis.pane.set_alpha(0.0)
    ax.yaxis.pane.set_alpha(0.0)
    ax.zaxis.pane.set_alpha(0.0)
    ax.grid(True, color=GRID, alpha=0.35)
    save(fig, "q2", "fig_q2_width_surface")


def fig_q3_layout() -> None:
    x, _, _, west, east = read_result3()
    x_min, x_max = -3704.0, 3704.0
    y_min, y_max = -1852.0, 1852.0
    bg_x = np.linspace(x_min, x_max, 350)
    bg = np.tile(np.linspace(1, 0, bg_x.size), (20, 1))

    fig, ax = plt.subplots(figsize=(8.4, 4.45), constrained_layout=True)
    ax.imshow(bg, extent=[x_min, x_max, y_min, y_max], origin="lower", cmap=DEPTH_CMAP, alpha=0.34, aspect="auto", zorder=0)
    ax.add_patch(Rectangle((x_min, y_min), x_max - x_min, y_max - y_min, fill=False, edgecolor=INK, lw=1.25, zorder=5))

    for wi, ei in zip(west, east):
        ax.add_patch(Rectangle((wi, y_min), ei - wi, y_max - y_min, facecolor="#9FD3C7", edgecolor="none", alpha=0.075, zorder=1))
        ax.plot([wi, wi], [y_min, y_max], color="#93B7BE", lw=0.35, alpha=0.55, zorder=2)
        ax.plot([ei, ei], [y_min, y_max], color="#93B7BE", lw=0.35, alpha=0.55, zorder=2)

    for i, xi in enumerate(x, 1):
        lw = 1.0 if i in (1, len(x)) else 0.65
        ax.plot([xi, xi], [y_min, y_max], color=NAVY, lw=lw, alpha=0.92, zorder=4)

    arrow2d(ax, (x_min + 260, y_min + 220), (x_min + 260, y_min + 680), color=INK, text="N", text_offset=(38, 8))
    arrow2d(ax, (x_min + 700, y_max + 130), (x_max - 700, y_max + 130), color=CLAY, text="由深到浅", text_offset=(-80, 26))
    ax.text(x_min + 120, y_max + 65, "西侧深水", color=CLAY, fontsize=10)
    ax.text(x_max - 860, y_max + 65, "东侧浅水", color=CLAY, fontsize=10)
    ax.text(x[3] + 36, y_max - 300, "测线", color=NAVY, fontsize=8)
    ax.text(west[5] - 145, y_min + 260, "覆盖边界", color="#71959A", fontsize=8)
    ax.set_xlim(x_min - 320, x_max + 320)
    ax.set_ylim(y_min - 260, y_max + 300)
    ax.set_xlabel("东西向坐标 x / m")
    ax.set_ylabel("南北向坐标 y / m")
    ax.set_title("规则坡面测线布设与覆盖条带")
    clean_axis(ax, grid=False)
    save(fig, "q3", "fig_q3_layout")


def fig_q3_spacing() -> None:
    x, spacing, eta, _, _ = read_result3()
    pair = np.arange(2, len(x) + 1)
    spacing = spacing[1:]
    eta = eta[1:]

    fig, axes = plt.subplots(1, 2, figsize=(9.8, 3.45), constrained_layout=True)
    axes[0].plot(pair, spacing, marker="o", color=BLUE, lw=1.95)
    axes[0].fill_between(pair, spacing, np.nanmin(spacing) - 35, color=BLUE, alpha=0.12)
    axes[0].set_xlabel("相邻测线对编号")
    axes[0].set_ylabel("测线间距 d / m")
    axes[0].set_title("最大可行间距递推")

    axes[1].axhspan(10, 20, color=TEAL, alpha=0.13)
    axes[1].plot(pair, eta, marker="o", color=TEAL, lw=1.95)
    axes[1].axhline(10, color=TEAL, ls="--", lw=1.0)
    axes[1].axhline(20, color=TEAL, ls="--", lw=1.0)
    axes[1].set_xlabel("相邻测线对编号")
    axes[1].set_ylabel("重叠率 eta / %")
    axes[1].set_ylim(8.8, 20.8)
    axes[1].text(4, 18.6, "约束区间 10%-20%", color=TEAL, fontsize=9)
    axes[1].set_title("相邻条带重叠率校验")
    for ax in axes:
        clean_axis(ax)
    save(fig, "q3", "fig_q3_spacing")


def plot_segments(ax, segs, color="white", lw=0.32, alpha=0.72, zorder=4):
    for x0, y0, y1 in segs:
        ax.plot([x0, x0], [y0, y1], color=color, lw=lw, alpha=alpha, zorder=zorder)


def fig_q4_depth() -> None:
    x, y, z = read_bathymetry()
    segs = read_result4_segments()
    x_grid, y_grid = np.meshgrid(x, y)

    fig, ax = plt.subplots(figsize=(7.35, 5.6), constrained_layout=True)
    levels = np.linspace(np.nanmin(z), np.nanmax(z), 18)
    im = ax.contourf(x_grid, y_grid, z, levels=levels, cmap=DEPTH_CMAP, extend="both")
    ax.contour(x_grid, y_grid, z, levels=np.arange(30, 201, 20), colors="#1F2D3A", linewidths=0.34, alpha=0.45)
    plot_segments(ax, segs, color="white", lw=0.36, alpha=0.82)
    ax.add_patch(Rectangle((x[0], y[0]), x[-1] - x[0], y[-1] - y[0], fill=False, edgecolor=INK, lw=1.05, zorder=5))
    cb = fig.colorbar(im, ax=ax, pad=0.016, shrink=0.94)
    cb.set_label("海水深度 / m")
    cb.set_ticks(np.arange(20, 201, 40))
    ax.set_xlabel("横向坐标 x / 海里")
    ax.set_ylabel("纵向坐标 y / 海里")
    ax.set_title("真实水深与推荐测线叠加")
    ax.text(0.02, 0.98, "0.50 NM 分带测线", transform=ax.transAxes, color="white", fontsize=8, weight="bold", va="top")
    clean_axis(ax, grid=False)
    save(fig, "q4", "fig_q4_depth")


def fig_q4_contour_route_overlay() -> None:
    x, y, z = read_bathymetry()
    segs = read_result4_segments()
    x_grid, y_grid = np.meshgrid(x, y)

    fig, ax = plt.subplots(figsize=(7.35, 5.55), constrained_layout=True)
    ax.contourf(x_grid, y_grid, z, levels=np.linspace(z.min(), z.max(), 12), cmap=DEPTH_CMAP, alpha=0.28)
    cs = ax.contour(x_grid, y_grid, z, levels=np.arange(20, 201, 20), colors=INK, linewidths=0.62, alpha=0.72)
    ax.clabel(cs, fmt="%d", fontsize=7, inline=True)
    plot_segments(ax, segs, color=BLUE, lw=0.38, alpha=0.84, zorder=5)
    for yy in np.arange(0.5, 5.0, 0.5):
        ax.axhline(yy, color=ACCENT, lw=0.55, alpha=0.35, zorder=3)
    ax.add_patch(Rectangle((x[0], y[0]), x[-1] - x[0], y[-1] - y[0], fill=False, edgecolor=INK, lw=1.05, zorder=6))
    ax.set_xlabel("横向坐标 x / 海里")
    ax.set_ylabel("纵向坐标 y / 海里")
    ax.set_title("等深线、分带边界与测线方案")
    ax.text(0.83, 0.97, "测线", transform=ax.transAxes, color=BLUE, fontsize=8, ha="left", va="top")
    ax.text(0.83, 0.92, "分带边界", transform=ax.transAxes, color=ACCENT, fontsize=8, ha="left", va="top")
    clean_axis(ax, grid=False)
    save(fig, "q4", "fig_q4_contour_route_overlay")


def fig_q4_bathymetry_surface() -> None:
    x, y, z = read_bathymetry()
    step = 4
    xs = x[::step]
    ys = y[::step]
    zs = z[::step, ::step]
    x_grid, y_grid = np.meshgrid(xs, ys)
    seabed = -zs

    fig = plt.figure(figsize=(7.6, 5.35))
    ax = fig.add_subplot(111, projection="3d")
    norm = Normalize(vmin=float(zs.min()), vmax=float(zs.max()))
    facecolors = DEPTH_CMAP(norm(zs))
    ax.plot_surface(x_grid, y_grid, seabed, facecolors=facecolors, linewidth=0, antialiased=True, alpha=0.97, shade=False)
    ax.contour(x_grid, y_grid, seabed, zdir="z", offset=-210, levels=12, cmap=DEPTH_CMAP, linewidths=0.78)
    mappable = ScalarMappable(norm=norm, cmap=DEPTH_CMAP)
    mappable.set_array(zs)
    cb = fig.colorbar(mappable, ax=ax, shrink=0.62, pad=0.08, label="海水深度 / m")
    cb.set_ticks(np.arange(40, 201, 40))
    ax.set_xlabel("x / 海里")
    ax.set_ylabel("y / 海里")
    ax.set_zlabel("海底高程 -D / m")
    ax.set_zlim(-210, -10)
    ax.set_title("附件真实海域三维海底地形")
    ax.view_init(elev=31, azim=-132)
    ax.set_box_aspect((1.05, 1.25, 0.62))
    ax.xaxis.pane.set_alpha(0.0)
    ax.yaxis.pane.set_alpha(0.0)
    ax.zaxis.pane.set_alpha(0.0)
    ax.grid(True, color=GRID, alpha=0.35)
    save(fig, "q4", "fig_q4_bathymetry_surface")


def fig_q4_comparison() -> None:
    data = read_result4_comparison()
    names = data["names"]
    length = data["length"]
    over20 = data["over20"]
    segments = data["segments"]
    max_eta = data["max_eta"]
    avg_eta = data["avg_eta"]

    fig, axes = plt.subplots(1, 2, figsize=(10.2, 4.55), constrained_layout=True, gridspec_kw={"width_ratios": [1.28, 1.0]})
    sizes = 45 + (segments - segments.min()) / (segments.max() - segments.min()) * 520
    colors = [MUTED, BLUE, AMBER, TEAL]
    lx = length / 1000
    oy = over20 / 1000

    ax = axes[0]
    focus = [1, 2, 3]
    ax.scatter(lx[focus], oy[focus], s=sizes[focus], c=[colors[i] for i in focus], alpha=0.86, edgecolor="white", linewidth=1.2, zorder=3)
    offsets = {1: (1.8, 3.5), 2: (3.0, 9.0), 3: (1.8, 8.5)}
    for i in focus:
        dx, dy = offsets[i]
        ax.text(lx[i] + dx, oy[i] + dy, names[i], fontsize=9.5, color=INK)
        ax.text(lx[i] + dx, oy[i] + dy - 5.8, f"{int(segments[i])} 段 | max {max_eta[i]:.1f}%", fontsize=7.6, color=MUTED)
    ax.annotate(
        "指标最优",
        xy=(lx[3], oy[3]),
        xytext=(421, 29),
        arrowprops=dict(arrowstyle="->", color=TEAL, lw=1.1),
        color=TEAL,
        fontsize=9,
    )
    ax.annotate(
        "工程折中推荐",
        xy=(lx[2], oy[2]),
        xytext=(438, 40),
        arrowprops=dict(arrowstyle="->", color=AMBER, lw=1.1),
        color=AMBER,
        fontsize=9,
    )
    ax.set_xlabel("测线总长度 / km")
    ax.set_ylabel("超 20% 重叠长度 / km")
    ax.set_title("关键分带方案局部权衡")
    ax.text(412.5, 55.5, f"0.50 NM 平均重叠率 {avg_eta[2]:.2f}%", color=MUTED, fontsize=8)
    ax.set_xlim(410, 456)
    ax.set_ylim(-8, 61)
    clean_axis(ax)

    ax2 = axes[1]
    ax2.scatter(lx, oy, s=sizes, c=colors, alpha=0.82, edgecolor="white", linewidth=1.2, zorder=3)
    overview_offsets = [(5, -15), (4, 8), (6, 0), (6, 12)]
    for i, name in enumerate(names):
        if i not in (0, 1):
            continue
        dx, dy = overview_offsets[i]
        ax2.text(lx[i] + dx, oy[i] + dy, name, fontsize=8.5, color=INK)
    ax2.text(418, 23, "低冗余方案见左图", fontsize=8, color=MUTED)
    ax2.text(468, 245, "圆面积表示线段数", color=MUTED, fontsize=8)
    ax2.set_xlabel("测线总长度 / km")
    ax2.set_ylabel("超 20% 重叠长度 / km")
    ax2.set_title("全局尺度概览")
    ax2.set_xlim(405, 575)
    ax2.set_ylim(-12, 305)
    clean_axis(ax2)
    save(fig, "q4", "fig_q4_comparison")


def main() -> None:
    fig_q1_geometry()
    fig_q1_result()
    fig_q2_3d_geometry()
    fig_q2_heatmap()
    fig_q2_width_surface()
    fig_q3_layout()
    fig_q3_spacing()
    fig_q4_depth()
    fig_q4_contour_route_overlay()
    fig_q4_bathymetry_surface()
    fig_q4_comparison()


if __name__ == "__main__":
    main()
