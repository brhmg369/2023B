"""Question 1 multibeam swath-width model for CUMCM 2023B.

The script computes the depth, swath coverage width, and overlap rate
for the locations requested in Table 1, then writes result1.xlsx.
"""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
RESULT_XLSX = ROOT / "result1.xlsx"
RESULT_CSV = ROOT / "docs" / "q1_result_table.csv"

THETA_DEG = 120.0
ALPHA_DEG = 1.5
CENTER_DEPTH_M = 70.0
POSITIONS_M = [-800, -600, -400, -200, 0, 200, 400, 600, 800]


@dataclass(frozen=True)
class Row:
    x_m: float
    depth_m: float
    left_half_width_m: float
    right_half_width_m: float
    coverage_width_m: float
    overlap_width_m: float | None
    overlap_rate_pct: float | None
    simplified_overlap_rate_pct: float | None


def depth_at(x_m: float, center_depth_m: float, alpha_rad: float) -> float:
    """Depth at horizontal offset x; positive x is the shallower upslope side."""
    return center_depth_m - x_m * math.tan(alpha_rad)


def half_widths(depth_m: float, theta_rad: float, alpha_rad: float) -> tuple[float, float]:
    """Return horizontal left and right half-widths of one swath strip.

    The seabed is approximated by a straight slope in the cross-track plane.
    The left half is toward the deeper side when x is ordered from negative to
    positive. Both widths are horizontal projections, so they are comparable
    with the horizontal line spacing in the table.
    """
    half_angle = theta_rad / 2.0
    common = depth_m * math.sin(half_angle) * math.cos(alpha_rad)
    left = common / math.cos(half_angle + alpha_rad)
    right = common / math.cos(half_angle - alpha_rad)
    return left, right


def compute_rows() -> list[Row]:
    theta_rad = math.radians(THETA_DEG)
    alpha_rad = math.radians(ALPHA_DEG)

    rows: list[Row] = []
    previous: Row | None = None
    for x_m in POSITIONS_M:
        depth_m = depth_at(x_m, CENTER_DEPTH_M, alpha_rad)
        left_half, right_half = half_widths(depth_m, theta_rad, alpha_rad)
        coverage = left_half + right_half

        overlap_width = None
        overlap_rate = None
        simplified_overlap_rate = None
        if previous is not None:
            spacing = x_m - previous.x_m
            overlap_width = previous.right_half_width_m + left_half - spacing
            overlap_rate = overlap_width / coverage * 100.0
            simplified_overlap_rate = (1.0 - spacing / coverage) * 100.0

        rows.append(
            Row(
                x_m=x_m,
                depth_m=depth_m,
                left_half_width_m=left_half,
                right_half_width_m=right_half,
                coverage_width_m=coverage,
                overlap_width_m=overlap_width,
                overlap_rate_pct=overlap_rate,
                simplified_overlap_rate_pct=simplified_overlap_rate,
            )
        )
        previous = rows[-1]

    return rows


def assert_flat_limit() -> None:
    """Check the model reduces to the textbook flat-bottom formula."""
    theta_rad = math.radians(THETA_DEG)
    flat_left, flat_right = half_widths(CENTER_DEPTH_M, theta_rad, 0.0)
    flat_width = flat_left + flat_right
    expected = 2.0 * CENTER_DEPTH_M * math.tan(theta_rad / 2.0)
    if not math.isclose(flat_width, expected, rel_tol=1e-12, abs_tol=1e-12):
        raise AssertionError((flat_width, expected))


def write_csv(rows: list[Row]) -> None:
    RESULT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with RESULT_CSV.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["指标"] + [int(row.x_m) for row in rows])
        writer.writerow(["海水深度/m"] + [f"{row.depth_m:.2f}" for row in rows])
        writer.writerow(["覆盖宽度/m"] + [f"{row.coverage_width_m:.2f}" for row in rows])
        writer.writerow(
            ["与前一条测线的重叠率/%"]
            + ["--" if row.overlap_rate_pct is None else f"{row.overlap_rate_pct:.2f}" for row in rows]
        )


def write_workbook(rows: list[Row]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "问题1结果"

    table = [
        ["测线距中心点处的距离/m"] + [int(row.x_m) for row in rows],
        ["海水深度/m"] + [round(row.depth_m, 2) for row in rows],
        ["覆盖宽度/m"] + [round(row.coverage_width_m, 2) for row in rows],
        [
            "与前一条测线的重叠率/%"
        ]
        + ["--" if row.overlap_rate_pct is None else round(row.overlap_rate_pct, 2) for row in rows],
    ]

    for r_idx, row_values in enumerate(table, start=1):
        for c_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if r_idx == 1 or c_idx == 1:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")

    ws.column_dimensions["A"].width = 28
    for c_idx in range(2, len(rows) + 2):
        ws.column_dimensions[get_column_letter(c_idx)].width = 12

    detail = wb.create_sheet("计算明细")
    headers = [
        "x/m",
        "D/m",
        "左半宽/m",
        "右半宽/m",
        "覆盖宽度/m",
        "重叠宽度/m",
        "精确重叠率/%",
        "平底简化重叠率/%(备用)",
    ]
    detail.append(headers)
    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        detail.append(
            [
                row.x_m,
                row.depth_m,
                row.left_half_width_m,
                row.right_half_width_m,
                row.coverage_width_m,
                row.overlap_width_m,
                row.overlap_rate_pct,
                row.simplified_overlap_rate_pct,
            ]
        )

    for column in detail.columns:
        detail.column_dimensions[column[0].column_letter].width = 18
        for cell in column:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    wb.save(RESULT_XLSX)


def print_table(rows: list[Row]) -> None:
    print("Table 1 results, exact overlap-rate definition")
    print(["指标"] + [int(row.x_m) for row in rows])
    print(["海水深度/m"] + [round(row.depth_m, 2) for row in rows])
    print(["覆盖宽度/m"] + [round(row.coverage_width_m, 2) for row in rows])
    print(
        ["与前一条测线的重叠率/%"]
        + ["--" if row.overlap_rate_pct is None else round(row.overlap_rate_pct, 2) for row in rows]
    )


def main() -> None:
    assert_flat_limit()
    rows = compute_rows()
    write_csv(rows)
    write_workbook(rows)
    print_table(rows)
    print(f"Wrote {RESULT_XLSX}")
    print(f"Wrote {RESULT_CSV}")


if __name__ == "__main__":
    main()
